import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import json
import os
import io
import time
import requests
from datetime import datetime
from supabase import create_client, Client

# ── CONFIG ────────────────────────────────────────────────────────────────────
SUPABASE_URL = "https://ujqhwlntfwezubuvgzeq.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InVqcWh3bG50ZndlenVidXZnemVxIiwicm9sZSI6ImFub24iLCJpYXQiOjE3Nzg2ODM0MTYsImV4cCI6MjA5NDI1OTQxNn0.Xwzub563z2FLA9PhFNqq-7Gu2xuscmd0xedcaQNb7js"
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")
GROQ_MODEL   = "llama-3.3-70b-versatile"

st.set_page_config(
    page_title="NABBR AI",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── STYLES ────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Bebas+Neue&family=DM+Sans:ital,wght@0,300;0,400;0,500;0,600;1,400&family=JetBrains+Mono:wght@400;500&display=swap');

:root {
  --bg:      #07080F;
  --bg2:     #0C0E1A;
  --bg3:     #111420;
  --bg4:     #161928;
  --border:  #1E2235;
  --border2: #252A40;
  --text:    #E8E9F5;
  --text2:   #7880A0;
  --text3:   #3A4060;
  --accent:  #4361EE;
  --accent2: #7B2FFF;
  --green:   #00D4AA;
  --red:     #FF4757;
  --yellow:  #FFD32A;
  --orange:  #FF6B35;
}

/* GLOBAL */
html,body,[class*="css"]{font-family:'DM Sans',sans-serif;background:var(--bg);color:var(--text);}
.main{background:var(--bg);}
.block-container{padding:2rem 2.5rem 4rem;}
h1,h2,h3{font-family:'Bebas Neue',sans-serif;letter-spacing:1px;}

/* HIDE STREAMLIT CHROME */
#MainMenu,footer,header{visibility:hidden;}
.stDeployButton{display:none;}

/* SIDEBAR */
[data-testid="stSidebar"]{background:var(--bg2);border-right:1px solid var(--border);}
[data-testid="stSidebar"] .block-container{padding:1.5rem 1rem;}

/* INPUTS */
.stTextInput>div>div>input,.stTextArea>div>div>textarea,.stSelectbox>div>div{
  background:var(--bg3)!important;border:1px solid var(--border2)!important;
  color:var(--text)!important;border-radius:8px!important;font-family:'DM Sans',sans-serif!important;}
.stTextInput>div>div>input:focus,.stTextArea>div>div>textarea:focus{border-color:var(--accent)!important;box-shadow:0 0 0 2px #4361EE20!important;}

/* BUTTONS */
.stButton>button{background:var(--accent);color:#fff;border:none;border-radius:8px;
  font-family:'DM Sans',sans-serif;font-weight:600;font-size:14px;padding:10px 24px;
  transition:all .2s;letter-spacing:.02em;}
.stButton>button:hover{background:#5472F0;transform:translateY(-1px);box-shadow:0 4px 20px #4361EE40;}

/* METRICS */
[data-testid="metric-container"]{background:var(--bg3);border:1px solid var(--border);
  border-radius:12px;padding:16px!important;}
[data-testid="metric-container"] label{color:var(--text2)!important;font-size:11px!important;
  font-weight:600;letter-spacing:.1em;text-transform:uppercase;}
[data-testid="metric-container"] [data-testid="stMetricValue"]{color:var(--text)!important;
  font-family:'Bebas Neue',sans-serif;font-size:28px!important;letter-spacing:1px;}

/* DATAFRAME */
[data-testid="stDataFrame"]{border:1px solid var(--border);border-radius:10px;overflow:hidden;}

/* TABS */
.stTabs [data-baseweb="tab-list"]{background:var(--bg3);border-radius:10px;padding:4px;gap:2px;border:1px solid var(--border);}
.stTabs [data-baseweb="tab"]{background:transparent;color:var(--text2);border-radius:8px;
  font-size:13px;font-weight:500;padding:8px 18px;}
.stTabs [aria-selected="true"]{background:var(--accent)!important;color:#fff!important;}

/* EXPANDER */
[data-testid="stExpander"]{background:var(--bg3);border:1px solid var(--border);border-radius:12px;}

/* ALERTS */
.stSuccess{background:#00D4AA15;border:1px solid #00D4AA40;color:var(--text);border-radius:10px;}
.stError{background:#FF475715;border:1px solid #FF475740;color:var(--text);border-radius:10px;}
.stWarning{background:#FFD32A15;border:1px solid #FFD32A40;color:var(--text);border-radius:10px;}
.stInfo{background:#4361EE15;border:1px solid #4361EE40;color:var(--text);border-radius:10px;}

/* FILE UPLOADER */
[data-testid="stFileUploader"]{background:var(--bg3);border:2px dashed var(--border2);
  border-radius:12px;padding:20px;}

/* PROGRESS */
.stProgress>div>div{background:var(--accent)!important;border-radius:4px;}

/* CUSTOM CARDS */
.nabbr-card{background:var(--bg3);border:1px solid var(--border);border-radius:14px;padding:20px;margin-bottom:14px;}
.nabbr-card:hover{border-color:var(--border2);}
.nabbr-hero{background:linear-gradient(135deg,var(--bg2) 0%,var(--bg3) 100%);
  border:1px solid var(--border);border-radius:16px;padding:32px;margin-bottom:24px;
  position:relative;overflow:hidden;}
.nabbr-hero::before{content:'';position:absolute;top:-40px;right:-40px;width:200px;height:200px;
  background:var(--accent);opacity:.04;border-radius:50%;filter:blur(40px);}
.agent-card{background:var(--bg3);border:1px solid var(--border);border-radius:12px;
  padding:16px;margin-bottom:10px;transition:all .2s;}
.agent-card.running{border-color:var(--accent);background:#4361EE08;}
.agent-card.done{border-color:var(--green);background:#00D4AA08;}
.agent-card.waiting{opacity:.5;}
.stat-pill{display:inline-block;background:var(--bg4);border:1px solid var(--border);
  border-radius:100px;padding:3px 12px;font-size:11px;color:var(--text2);font-weight:500;
  font-family:'JetBrains Mono',monospace;margin:2px;}
.badge{display:inline-block;padding:3px 10px;border-radius:100px;font-size:10px;
  font-weight:700;letter-spacing:.08em;text-transform:uppercase;}
.badge-blue{background:#4361EE20;color:#4361EE;border:1px solid #4361EE40;}
.badge-green{background:#00D4AA20;color:#00D4AA;border:1px solid #00D4AA40;}
.badge-red{background:#FF475720;color:#FF4757;border:1px solid #FF475740;}
.badge-yellow{background:#FFD32A20;color:#FFD32A;border:1px solid #FFD32A40;}
.badge-purple{background:#7B2FFF20;color:#7B2FFF;border:1px solid #7B2FFF40;}
.nabbr-logo{font-family:'Bebas Neue',sans-serif;font-size:28px;letter-spacing:3px;
  background:linear-gradient(135deg,#4361EE,#7B2FFF);-webkit-background-clip:text;
  -webkit-text-fill-color:transparent;background-clip:text;}
.section-title{font-family:'Bebas Neue',sans-serif;font-size:22px;letter-spacing:2px;
  color:var(--text);margin-bottom:16px;padding-bottom:8px;
  border-bottom:1px solid var(--border);}
.insight-item{background:var(--bg4);border-left:3px solid var(--accent);
  border-radius:0 8px 8px 0;padding:10px 14px;margin-bottom:8px;font-size:13px;
  color:var(--text2);line-height:1.6;}
.insight-item strong{color:var(--text);}
</style>
""", unsafe_allow_html=True)

# ── SUPABASE CLIENT ───────────────────────────────────────────────────────────
@st.cache_resource
def get_supabase() -> Client:
    return create_client(SUPABASE_URL, SUPABASE_KEY)

supabase = get_supabase()

# ── SESSION STATE INIT ────────────────────────────────────────────────────────
def init_state():
    defaults = {
        "user": None, "session_token": None, "page": "login",
        "df": None, "filename": None, "cleaned_df": None,
        "analysis": {}, "workspace_id": None, "session_id": None,
        "agent_results": {}, "ml_results": None, "charts": [],
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v

init_state()

# ── GROQ AI CALL ──────────────────────────────────────────────────────────────
def ask_groq(system: str, user: str, max_tokens: int = 1024) -> str:
    if not GROQ_API_KEY:
        return "⚠️ GROQ_API_KEY not set. Add it to your environment variables."
    try:
        r = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"},
            json={"model": GROQ_MODEL, "max_tokens": max_tokens, "temperature": 0.4,
                  "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]},
            timeout=30,
        )
        data = r.json()
        return data["choices"][0]["message"]["content"]
    except Exception as e:
        return f"Error calling Groq: {str(e)}"

# ── DATA HELPERS ──────────────────────────────────────────────────────────────
def load_file(file) -> pd.DataFrame:
    name = file.name.lower()
    if name.endswith(".csv"):
        return pd.read_csv(file)
    elif name.endswith((".xlsx", ".xls")):
        return pd.read_excel(file)
    elif name.endswith(".json"):
        return pd.read_json(file)
    else:
        st.error("Unsupported file type. Please upload CSV, Excel, or JSON.")
        return None

def auto_clean(df: pd.DataFrame) -> tuple:
    original_shape = df.shape
    report = []
    df = df.copy()

    # Drop duplicates
    before = len(df)
    df = df.drop_duplicates()
    dropped_dups = before - len(df)
    if dropped_dups > 0:
        report.append(f"✅ Removed **{dropped_dups}** duplicate rows")

    # Handle missing values
    null_counts = df.isnull().sum()
    for col in df.columns:
        null_pct = df[col].isnull().mean()
        if null_pct > 0.5:
            df = df.drop(columns=[col])
            report.append(f"🗑️ Dropped column **{col}** ({null_pct:.0%} missing)")
        elif null_pct > 0:
            if df[col].dtype in ["float64", "int64"]:
                df[col] = df[col].fillna(df[col].median())
                report.append(f"📊 Filled **{col}** with median ({df[col].median():.2f})")
            else:
                df[col] = df[col].fillna(df[col].mode()[0] if len(df[col].mode()) > 0 else "Unknown")
                report.append(f"📝 Filled **{col}** with mode")

    # Remove outliers (IQR) for numeric columns
    numeric_cols = df.select_dtypes(include=[np.number]).columns
    outliers_removed = 0
    for col in numeric_cols:
        Q1, Q3 = df[col].quantile(0.25), df[col].quantile(0.75)
        IQR = Q3 - Q1
        before = len(df)
        df = df[~((df[col] < Q1 - 3 * IQR) | (df[col] > Q3 + 3 * IQR))]
        outliers_removed += before - len(df)
    if outliers_removed > 0:
        report.append(f"🎯 Removed **{outliers_removed}** outliers (3×IQR method)")

    # Parse date columns
    for col in df.columns:
        if any(kw in col.lower() for kw in ["date", "time", "dt", "created", "updated"]):
            try:
                df[col] = pd.to_datetime(df[col])
                report.append(f"📅 Parsed **{col}** as datetime")
            except:
                pass

    if not report:
        report.append("✅ Data is already clean — no issues found!")

    return df, report, original_shape

def get_data_profile(df: pd.DataFrame) -> dict:
    numeric = df.select_dtypes(include=[np.number])
    categorical = df.select_dtypes(include=["object", "category"])
    datetime_cols = df.select_dtypes(include=["datetime64"])
    return {
        "rows": len(df), "cols": len(df.columns),
        "numeric_cols": list(numeric.columns),
        "categorical_cols": list(categorical.columns),
        "datetime_cols": list(datetime_cols.columns),
        "null_count": int(df.isnull().sum().sum()),
        "memory_mb": round(df.memory_usage(deep=True).sum() / 1024**2, 2),
        "dtypes": df.dtypes.astype(str).to_dict(),
        "sample": df.head(3).to_dict(),
        "describe": numeric.describe().round(2).to_dict() if len(numeric.columns) > 0 else {},
        "col_nulls": df.isnull().sum().to_dict(),
    }

# ── AGENTS ────────────────────────────────────────────────────────────────────
AGENTS = [
    {"id": "inspector",   "name": "Quality Inspector",     "icon": "🔍", "color": "#4361EE"},
    {"id": "statistician","name": "Statistical Analyst",   "icon": "📊", "color": "#7B2FFF"},
    {"id": "ml_expert",   "name": "ML Expert",             "icon": "🧠", "color": "#00D4AA"},
    {"id": "visualizer",  "name": "Visualization Designer","icon": "🎨", "color": "#FF6B35"},
    {"id": "reporter",    "name": "Report Writer",         "icon": "📄", "color": "#FFD32A"},
    {"id": "advisor",     "name": "Business Advisor",      "icon": "💼", "color": "#FF4757"},
    {"id": "strategist",  "name": "Strategy Director",     "icon": "🎯", "color": "#4361EE"},
]

def run_agent(agent_id: str, profile: dict, df: pd.DataFrame) -> str:
    df_info = f"""
Dataset: {profile['rows']} rows × {profile['cols']} columns
Numeric columns: {', '.join(profile['numeric_cols'][:10])}
Categorical columns: {', '.join(profile['categorical_cols'][:10])}
Datetime columns: {', '.join(profile['datetime_cols'])}
Missing values: {profile['null_count']}
Statistics: {json.dumps(profile['describe'], default=str)[:1000]}
"""
    prompts = {
        "inspector": (
            "You are a Senior Data Quality Inspector. Analyze datasets for quality issues.",
            f"Analyze this dataset for data quality:\n{df_info}\n\nProvide: 1) Overall quality score (0-100), 2) Top 5 quality issues found, 3) Data completeness assessment, 4) Consistency checks, 5) Recommendations to improve quality. Be specific with column names and numbers."
        ),
        "statistician": (
            "You are a Senior Statistical Analyst. Provide deep statistical insights.",
            f"Perform statistical analysis on this dataset:\n{df_info}\n\nProvide: 1) Key statistical findings, 2) Distribution patterns, 3) Correlations and relationships, 4) Anomalies detected, 5) Statistical significance of patterns. Use specific numbers from the data."
        ),
        "ml_expert": (
            "You are an ML Expert. Recommend and explain ML approaches for datasets.",
            f"Analyze this dataset for ML opportunities:\n{df_info}\n\nProvide: 1) Best ML task (classification/regression/clustering/forecasting), 2) Recommended target column and why, 3) Top 5 most predictive features, 4) Best algorithms to try, 5) Expected performance range. Be specific."
        ),
        "visualizer": (
            "You are a Data Visualization Designer. Recommend the best charts for insights.",
            f"Recommend visualizations for this dataset:\n{df_info}\n\nProvide: 1) Top 5 most insightful charts to create, 2) What each chart will reveal, 3) Color and layout recommendations, 4) Dashboard layout suggestion, 5) Key story the data tells visually."
        ),
        "reporter": (
            "You are a Business Report Writer. Create executive-ready reports from data.",
            f"Write an executive summary for this dataset:\n{df_info}\n\nProvide: 1) Executive summary (3-4 sentences), 2) Key findings (top 5 bullet points), 3) Data highlights with specific numbers, 4) Areas of concern, 5) Recommended next steps for the business."
        ),
        "advisor": (
            "You are a Senior Business Advisor. Translate data into business strategy.",
            f"Provide business advisory insights for this dataset:\n{df_info}\n\nProvide: 1) Business context interpretation, 2) Top 3 business opportunities identified, 3) Top 3 risks or concerns, 4) Competitive implications, 5) Specific actionable recommendations with expected ROI."
        ),
        "strategist": (
            "You are a Strategy Director. Create 90-day strategic plans from data insights.",
            f"Create a strategic plan based on this dataset:\n{df_info}\n\nProvide: 1) Strategic situation assessment, 2) 30-day quick wins (3 actions), 3) 60-day growth initiatives (3 actions), 4) 90-day transformation goals (3 actions), 5) KPIs to track success. Be specific and actionable."
        ),
    }
    system, user = prompts[agent_id]
    return ask_groq(system, user, max_tokens=800)

# ── ML PIPELINE ───────────────────────────────────────────────────────────────
def run_ml_pipeline(df: pd.DataFrame, target_col: str, task: str) -> dict:
    from sklearn.model_selection import train_test_split
    from sklearn.preprocessing import LabelEncoder, StandardScaler
    from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor, GradientBoostingClassifier, GradientBoostingRegressor
    from sklearn.linear_model import LogisticRegression, LinearRegression
    from sklearn.metrics import (accuracy_score, f1_score, r2_score,
                                  mean_absolute_error, classification_report, confusion_matrix)
    from sklearn.impute import SimpleImputer

    results = {}
    df_ml = df.copy()

    # Prepare features
    feature_cols = [c for c in df_ml.columns if c != target_col]
    X = df_ml[feature_cols]
    y = df_ml[target_col]

    # Encode categoricals
    le_dict = {}
    for col in X.select_dtypes(include=["object", "category"]).columns:
        le = LabelEncoder()
        X[col] = le.fit_transform(X[col].astype(str))
        le_dict[col] = le

    # Drop datetime
    X = X.select_dtypes(exclude=["datetime64"])

    # Impute
    imp = SimpleImputer(strategy="median")
    X_imp = imp.fit_transform(X)

    # Encode target if classification
    if task == "classification":
        le_y = LabelEncoder()
        y = le_y.fit_transform(y.astype(str))
        results["classes"] = list(le_y.classes_)

    X_train, X_test, y_train, y_test = train_test_split(X_imp, y, test_size=0.2, random_state=42)

    # Scale
    scaler = StandardScaler()
    X_train_s = scaler.fit_transform(X_train)
    X_test_s  = scaler.transform(X_test)

    if task == "classification":
        models = {
            "Random Forest":       RandomForestClassifier(n_estimators=100, random_state=42),
            "Gradient Boosting":   GradientBoostingClassifier(n_estimators=100, random_state=42),
            "Logistic Regression": LogisticRegression(max_iter=1000, random_state=42),
        }
        best_model, best_score, best_name = None, 0, ""
        model_scores = {}
        for name, model in models.items():
            model.fit(X_train_s, y_train)
            preds = model.predict(X_test_s)
            score = accuracy_score(y_test, preds)
            f1    = f1_score(y_test, preds, average="weighted", zero_division=0)
            model_scores[name] = {"accuracy": round(score, 4), "f1": round(f1, 4)}
            if score > best_score:
                best_score, best_model, best_name = score, model, name

        preds = best_model.predict(X_test_s)
        results.update({
            "task": "classification", "best_model": best_name,
            "accuracy": round(best_score, 4),
            "f1": round(f1_score(y_test, preds, average="weighted", zero_division=0), 4),
            "model_scores": model_scores,
            "y_test": list(y_test[:100]), "y_pred": list(preds[:100]),
        })
        if hasattr(best_model, "feature_importances_"):
            fi = pd.Series(best_model.feature_importances_, index=X.columns)
            results["feature_importance"] = fi.sort_values(ascending=False).head(10).to_dict()

    else:  # regression
        models = {
            "Random Forest":     RandomForestRegressor(n_estimators=100, random_state=42),
            "Gradient Boosting": GradientBoostingRegressor(n_estimators=100, random_state=42),
            "Linear Regression": LinearRegression(),
        }
        best_model, best_score, best_name = None, -np.inf, ""
        model_scores = {}
        for name, model in models.items():
            model.fit(X_train_s, y_train)
            preds = model.predict(X_test_s)
            score = r2_score(y_test, preds)
            mae   = mean_absolute_error(y_test, preds)
            model_scores[name] = {"r2": round(score, 4), "mae": round(mae, 2)}
            if score > best_score:
                best_score, best_model, best_name = score, model, name

        preds = best_model.predict(X_test_s)
        results.update({
            "task": "regression", "best_model": best_name,
            "r2": round(best_score, 4),
            "mae": round(mean_absolute_error(y_test, preds), 2),
            "model_scores": model_scores,
            "y_test": list(y_test[:100]), "y_pred": list(preds[:100]),
        })
        if hasattr(best_model, "feature_importances_"):
            fi = pd.Series(best_model.feature_importances_, index=X.columns)
            results["feature_importance"] = fi.sort_values(ascending=False).head(10).to_dict()

    results["feature_cols"] = list(X.columns)
    results["n_train"] = len(X_train)
    results["n_test"]  = len(X_test)
    return results

# ── PDF REPORT ────────────────────────────────────────────────────────────────
def generate_pdf_report(filename, profile, agent_results, ml_results) -> bytes:
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        # Header
        pdf.set_fill_color(7, 8, 15)
        pdf.rect(0, 0, 210, 40, 'F')
        pdf.set_font("Helvetica", "B", 28)
        pdf.set_text_color(67, 97, 238)
        pdf.cell(0, 15, "", ln=True)
        pdf.cell(0, 10, "NABBR AI", ln=True, align="C")
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(120, 128, 160)
        pdf.cell(0, 6, "Enterprise Analytics Report — by Nalla", ln=True, align="C")
        pdf.set_text_color(0, 0, 0)
        pdf.ln(20)

        # Dataset info
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(30, 30, 30)
        pdf.cell(0, 8, f"Dataset: {filename}", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(80, 80, 80)
        pdf.cell(0, 6, f"Generated: {datetime.now().strftime('%B %d, %Y at %H:%M')}", ln=True)
        pdf.cell(0, 6, f"Rows: {profile['rows']:,}  |  Columns: {profile['cols']}  |  Missing Values: {profile['null_count']}", ln=True)
        pdf.ln(8)

        # Agent results
        for agent in AGENTS:
            aid = agent["id"]
            if aid not in agent_results:
                continue
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(67, 97, 238)
            pdf.cell(0, 8, f"{agent['icon']} {agent['name']}", ln=True)
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(50, 50, 50)
            text = agent_results[aid].replace("**", "").replace("*", "").replace("#", "")
            pdf.multi_cell(0, 5, text[:1500])
            pdf.ln(4)

        # ML Results
        if ml_results:
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(67, 97, 238)
            pdf.cell(0, 8, "ML Model Results", ln=True)
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(50, 50, 50)
            if ml_results.get("task") == "classification":
                pdf.cell(0, 6, f"Best Model: {ml_results['best_model']}  |  Accuracy: {ml_results['accuracy']:.2%}  |  F1: {ml_results['f1']:.4f}", ln=True)
            else:
                pdf.cell(0, 6, f"Best Model: {ml_results['best_model']}  |  R²: {ml_results['r2']:.4f}  |  MAE: {ml_results['mae']:,.2f}", ln=True)

        # Footer
        pdf.set_y(-20)
        pdf.set_font("Helvetica", "I", 8)
        pdf.set_text_color(150, 150, 150)
        pdf.cell(0, 6, "NABBR AI by Nalla — Confidential Analytics Report", align="C")

        return bytes(pdf.output())
    except Exception as e:
        return f"PDF generation error: {e}".encode()

# ── EXCEL REPORT ──────────────────────────────────────────────────────────────
def generate_excel_report(df, profile, agent_results, ml_results) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary
    ws1 = wb.active
    ws1.title = "Executive Summary"
    header_fill = PatternFill("solid", fgColor="4361EE")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    ws1["A1"] = "NABBR AI — Analytics Report by Nalla"
    ws1["A1"].font = Font(bold=True, size=16, color="4361EE")
    ws1["A2"] = f"Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}"
    ws1["A2"].font = Font(color="888888", italic=True)
    ws1["A4"] = "Dataset Profile"
    ws1["A4"].font = Font(bold=True, size=12)
    for i, (k, v) in enumerate([("Rows", profile["rows"]), ("Columns", profile["cols"]),
                                  ("Missing Values", profile["null_count"]), ("Memory (MB)", profile["memory_mb"])], start=5):
        ws1[f"A{i}"] = k
        ws1[f"B{i}"] = v
        ws1[f"A{i}"].font = Font(bold=True)
    row = 10
    for agent in AGENTS:
        if agent["id"] in agent_results:
            ws1[f"A{row}"] = f"{agent['icon']} {agent['name']}"
            ws1[f"A{row}"].font = Font(bold=True, color="4361EE")
            row += 1
            text = agent_results[agent["id"]].replace("**", "").replace("*", "")
            for line in text.split("\n")[:15]:
                ws1[f"A{row}"] = line.strip()
                ws1[f"A{row}"].alignment = Alignment(wrap_text=True)
                row += 1
            row += 1
    ws1.column_dimensions["A"].width = 40
    ws1.column_dimensions["B"].width = 20

    # ── Sheet 2: Raw Data
    ws2 = wb.create_sheet("Data")
    for r in dataframe_to_rows(df.head(1000), index=False, header=True):
        ws2.append(r)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font

    # ── Sheet 3: Statistics
    ws3 = wb.create_sheet("Statistics")
    numeric = df.select_dtypes(include=[np.number])
    if len(numeric.columns) > 0:
        stats_df = numeric.describe().round(2)
        for r in dataframe_to_rows(stats_df.reset_index(), index=False, header=True):
            ws3.append(r)
        for cell in ws3[1]:
            cell.fill = header_fill
            cell.font = header_font

    # ── Sheet 4: ML Results
    if ml_results:
        ws4 = wb.create_sheet("ML Results")
        ws4["A1"] = "ML Model Results"
        ws4["A1"].font = Font(bold=True, size=14, color="4361EE")
        ws4["A3"] = "Task"
        ws4["B3"] = ml_results.get("task", "")
        ws4["A4"] = "Best Model"
        ws4["B4"] = ml_results.get("best_model", "")
        if ml_results.get("task") == "classification":
            ws4["A5"] = "Accuracy"
            ws4["B5"] = f"{ml_results.get('accuracy', 0):.2%}"
            ws4["A6"] = "F1 Score"
            ws4["B6"] = f"{ml_results.get('f1', 0):.4f}"
        else:
            ws4["A5"] = "R² Score"
            ws4["B5"] = f"{ml_results.get('r2', 0):.4f}"
            ws4["A6"] = "MAE"
            ws4["B6"] = f"{ml_results.get('mae', 0):,.2f}"
        if "feature_importance" in ml_results:
            ws4["A8"] = "Feature Importance"
            ws4["A8"].font = Font(bold=True)
            for i, (feat, imp) in enumerate(ml_results["feature_importance"].items(), start=9):
                ws4[f"A{i}"] = feat
                ws4[f"B{i}"] = round(imp, 4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ── AUTH FUNCTIONS ────────────────────────────────────────────────────────────
def login(email: str, password: str):
    try:
        res = supabase.auth.sign_in_with_password({"email": email, "password": password})
        if res.user:
            st.session_state.user = res.user
            st.session_state.session_token = res.session.access_token
            st.session_state.page = "home"
            return True, "Login successful!"
        return False, "Invalid credentials"
    except Exception as e:
        return False, str(e)

def register(email: str, password: str, name: str):
    try:
        res = supabase.auth.sign_up({"email": email, "password": password, "options": {"data": {"full_name": name}}})
        if res.user:
            return True, "Account created! Please check your email to verify, then log in."
        return False, "Registration failed"
    except Exception as e:
        return False, str(e)

def logout():
    supabase.auth.sign_out()
    for key in list(st.session_state.keys()):
        del st.session_state[key]
    init_state()
    st.rerun()

# ── WORKSPACE FUNCTIONS ───────────────────────────────────────────────────────
def get_workspaces():
    try:
        user_id = st.session_state.user.id
        res = supabase.table("workspaces").select("*").eq("user_id", user_id).order("created_at", desc=True).execute()
        return res.data or []
    except:
        return []

def create_workspace(name: str, desc: str = ""):
    try:
        user_id = st.session_state.user.id
        res = supabase.table("workspaces").insert({"user_id": user_id, "name": name, "description": desc}).execute()
        return res.data[0] if res.data else None
    except Exception as e:
        st.error(f"Error creating workspace: {e}")
        return None

def get_sessions(workspace_id: str):
    try:
        res = supabase.table("sessions").select("*").eq("workspace_id", workspace_id).order("created_at", desc=True).execute()
        return res.data or []
    except:
        return []

def save_session(workspace_id, name, filename, profile, analysis):
    try:
        user_id = st.session_state.user.id
        res = supabase.table("sessions").insert({
            "workspace_id": workspace_id, "user_id": user_id,
            "name": name, "filename": filename,
            "row_count": profile["rows"], "col_count": profile["cols"],
            "analysis": analysis,
        }).execute()
        return res.data[0]["id"] if res.data else None
    except Exception as e:
        st.error(f"Error saving session: {e}")
        return None

# ── SIDEBAR ───────────────────────────────────────────────────────────────────
def render_sidebar():
    with st.sidebar:
        st.markdown('<div class="nabbr-logo">NABBR AI</div>', unsafe_allow_html=True)
        st.markdown('<p style="font-size:11px;color:#3A4060;font-family:\'JetBrains Mono\',monospace;margin-top:-4px;">by Nalla · Enterprise Analytics</p>', unsafe_allow_html=True)
        st.divider()

        if st.session_state.user:
            user_meta = st.session_state.user.user_metadata or {}
            name = user_meta.get("full_name", st.session_state.user.email.split("@")[0])
            st.markdown(f'<div style="background:#111420;border:1px solid #1E2235;border-radius:10px;padding:12px;margin-bottom:16px;"><span style="font-size:11px;color:#7880A0;">Logged in as</span><br><strong style="font-size:14px;">{name}</strong><br><span style="font-size:11px;color:#3A4060;">{st.session_state.user.email}</span></div>', unsafe_allow_html=True)

            pages = [
                ("🏠", "Home", "home"),
                ("📁", "Upload & Clean", "upload"),
                ("🤖", "AI Agents", "agents"),
                ("📊", "Dashboard", "dashboard"),
                ("🧠", "ML Pipeline", "ml"),
                ("📄", "Reports", "reports"),
                ("📜", "History", "history"),
            ]
            st.markdown("**Navigation**")
            for icon, label, page_id in pages:
                active = st.session_state.page == page_id
                if st.button(f"{icon} {label}", key=f"nav_{page_id}",
                             use_container_width=True,
                             type="primary" if active else "secondary"):
                    st.session_state.page = page_id
                    st.rerun()

            st.divider()
            if st.button("🚪 Logout", use_container_width=True, type="secondary"):
                logout()

# ── PAGES ─────────────────────────────────────────────────────────────────────

def page_login():
    col1, col2, col3 = st.columns([1, 1.2, 1])
    with col2:
        st.markdown('<br><br>', unsafe_allow_html=True)
        st.markdown('<div style="text-align:center"><span style="font-family:\'Bebas Neue\',sans-serif;font-size:52px;letter-spacing:4px;background:linear-gradient(135deg,#4361EE,#7B2FFF);-webkit-background-clip:text;-webkit-text-fill-color:transparent;background-clip:text;">NABBR AI</span><br><span style="font-size:12px;color:#3A4060;letter-spacing:.2em;font-family:\'JetBrains Mono\',monospace;">BY NALLA · ENTERPRISE ANALYTICS</span></div>', unsafe_allow_html=True)
        st.markdown('<p style="text-align:center;color:#7880A0;font-size:14px;margin:12px 0 28px;">From raw data to boardroom-ready insights — in seconds.</p>', unsafe_allow_html=True)

        tab_login, tab_reg = st.tabs(["Sign In", "Create Account"])

        with tab_login:
            st.markdown('<br>', unsafe_allow_html=True)
            email    = st.text_input("Email", placeholder="you@example.com", key="li_email")
            password = st.text_input("Password", type="password", placeholder="••••••••", key="li_pass")
            if st.button("Sign In →", use_container_width=True, type="primary"):
                if email and password:
                    with st.spinner("Signing in..."):
                        ok, msg = login(email, password)
                    if ok:
                        st.success(msg)
                        time.sleep(0.5)
                        st.rerun()
                    else:
                        st.error(msg)
                else:
                    st.warning("Please enter email and password.")

        with tab_reg:
            st.markdown('<br>', unsafe_allow_html=True)
            name     = st.text_input("Full Name", placeholder="Your name", key="reg_name")
            email    = st.text_input("Email", placeholder="you@example.com", key="reg_email")
            password = st.text_input("Password", type="password", placeholder="Min 6 characters", key="reg_pass")
            if st.button("Create Account →", use_container_width=True, type="primary"):
                if name and email and password:
                    with st.spinner("Creating account..."):
                        ok, msg = register(email, password, name)
                    if ok:
                        st.success(msg)
                    else:
                        st.error(msg)
                else:
                    st.warning("Please fill all fields.")

        st.markdown('<br><br><p style="text-align:center;font-size:11px;color:#3A4060;">⚡ Powered by LLaMA 3.3 · Built by Nalla</p>', unsafe_allow_html=True)


def page_home():
    user_meta = st.session_state.user.user_metadata or {}
    name = user_meta.get("full_name", "there")

    st.markdown(f"""
    <div class="nabbr-hero">
      <h1 style="font-family:'Bebas Neue',sans-serif;font-size:36px;letter-spacing:2px;margin:0;">
        Welcome back, {name.split()[0]} ⚡
      </h1>
      <p style="color:#7880A0;font-size:14px;margin:8px 0 0;">
        NABBR AI · Enterprise Analytics Platform by Nalla
      </p>
    </div>
    """, unsafe_allow_html=True)

    # Workspaces
    st.markdown('<div class="section-title">YOUR WORKSPACES</div>', unsafe_allow_html=True)
    workspaces = get_workspaces()

    col_new, col_space = st.columns([1, 3])
    with col_new:
        with st.expander("➕ New Workspace"):
            ws_name = st.text_input("Workspace Name", placeholder="e.g. Q4 Sales Analysis")
            ws_desc = st.text_area("Description", placeholder="What are you analyzing?", height=80)
            if st.button("Create Workspace", type="primary"):
                if ws_name:
                    ws = create_workspace(ws_name, ws_desc)
                    if ws:
                        st.success(f"✅ '{ws_name}' created!")
                        st.rerun()
                else:
                    st.warning("Please enter a workspace name.")

    if not workspaces:
        st.info("No workspaces yet. Create one above to get started!")
    else:
        cols = st.columns(3)
        for i, ws in enumerate(workspaces):
            with cols[i % 3]:
                sessions = get_sessions(ws["id"])
                created = ws["created_at"][:10] if ws.get("created_at") else ""
                st.markdown(f"""
                <div class="nabbr-card" style="cursor:pointer;">
                  <div style="display:flex;justify-content:space-between;align-items:flex-start;">
                    <div>
                      <strong style="font-size:16px;">{ws['name']}</strong>
                      <p style="color:#7880A0;font-size:12px;margin:4px 0;">{ws.get('description','')[:60]}</p>
                    </div>
                    <span class="badge badge-blue">{len(sessions)} sessions</span>
                  </div>
                  <p style="font-size:11px;color:#3A4060;font-family:'JetBrains Mono',monospace;margin-top:8px;">{created}</p>
                </div>
                """, unsafe_allow_html=True)
                if st.button(f"Open →", key=f"ws_{ws['id']}", use_container_width=True):
                    st.session_state.workspace_id = ws["id"]
                    st.session_state.page = "upload"
                    st.rerun()

    # Quick stats
    st.divider()
    st.markdown('<div class="section-title">PLATFORM OVERVIEW</div>', unsafe_allow_html=True)
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Workspaces", len(workspaces))
    total_sessions = sum(len(get_sessions(ws["id"])) for ws in workspaces)
    c2.metric("Total Sessions", total_sessions)
    c3.metric("AI Agents", "7")
    c4.metric("ML Models", "3 per run")


def page_upload():
    st.markdown('<div class="section-title">📁 DATA UPLOAD & AUTO-CLEAN</div>', unsafe_allow_html=True)

    # Workspace selector
    workspaces = get_workspaces()
    if not workspaces:
        st.warning("Please create a workspace first.")
        if st.button("Go to Home"):
            st.session_state.page = "home"
            st.rerun()
        return

    ws_names = {ws["name"]: ws["id"] for ws in workspaces}
    selected_ws = st.selectbox("Select Workspace", list(ws_names.keys()))
    st.session_state.workspace_id = ws_names[selected_ws]

    uploaded = st.file_uploader(
        "Drop your dataset here — CSV, Excel, or JSON",
        type=["csv", "xlsx", "xls", "json"],
        help="Max 200MB. Supports CSV, Excel (.xlsx/.xls), and JSON."
    )

    if uploaded:
        with st.spinner("Loading data..."):
            df = load_file(uploaded)

        if df is not None:
            st.session_state.df = df
            st.session_state.filename = uploaded.name

            # Preview
            st.markdown("### 👀 Raw Data Preview")
            col1, col2, col3, col4 = st.columns(4)
            col1.metric("Rows", f"{len(df):,}")
            col2.metric("Columns", len(df.columns))
            col3.metric("Missing Values", int(df.isnull().sum().sum()))
            col4.metric("Memory", f"{df.memory_usage(deep=True).sum()/1024**2:.1f} MB")
            st.dataframe(df.head(10), use_container_width=True)

            # Auto clean
            st.markdown("### 🧹 Auto-Clean Pipeline")
            if st.button("⚡ Run Auto-Clean", type="primary"):
                with st.spinner("Cleaning data..."):
                    cleaned_df, report, orig_shape = auto_clean(df)
                    st.session_state.cleaned_df = cleaned_df

                st.success(f"✅ Cleaning complete! {orig_shape[0]:,} → {len(cleaned_df):,} rows")
                for item in report:
                    st.markdown(f'<div class="insight-item">{item}</div>', unsafe_allow_html=True)

                st.markdown("### ✅ Cleaned Data Preview")
                col1, col2, col3 = st.columns(3)
                col1.metric("Rows After", f"{len(cleaned_df):,}", f"{len(cleaned_df)-orig_shape[0]:,}")
                col2.metric("Columns After", len(cleaned_df.columns))
                col3.metric("Missing Values", int(cleaned_df.isnull().sum().sum()))
                st.dataframe(cleaned_df.head(10), use_container_width=True)

                # Save session
                session_name = f"{uploaded.name} — {datetime.now().strftime('%b %d %H:%M')}"
                profile = get_data_profile(cleaned_df)
                session_id = save_session(
                    st.session_state.workspace_id,
                    session_name, uploaded.name, profile, {}
                )
                if session_id:
                    st.session_state.session_id = session_id
                    st.info(f"💾 Session saved: '{session_name}'")

                if st.button("🤖 Run AI Agents →", type="primary"):
                    st.session_state.page = "agents"
                    st.rerun()

    elif st.session_state.df is not None:
        st.info(f"✅ Dataset loaded: **{st.session_state.filename}** ({len(st.session_state.df):,} rows × {len(st.session_state.df.columns)} cols)")
        if st.button("🤖 Go to AI Agents →", type="primary"):
            st.session_state.page = "agents"
            st.rerun()


def page_agents():
    st.markdown('<div class="section-title">🤖 AI AGENT ANALYSIS</div>', unsafe_allow_html=True)

    df = st.session_state.cleaned_df if st.session_state.cleaned_df is not None else st.session_state.df
    if df is None:
        st.warning("Please upload a dataset first.")
        if st.button("Upload Data"):
            st.session_state.page = "upload"
            st.rerun()
        return

    profile = get_data_profile(df)

    # Show data summary
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Rows", f"{profile['rows']:,}")
    col2.metric("Columns", profile["cols"])
    col3.metric("Numeric Cols", len(profile["numeric_cols"]))
    col4.metric("Categorical Cols", len(profile["categorical_cols"]))

    st.markdown('<br>', unsafe_allow_html=True)

    # Agent selector
    st.markdown("### Select Agents to Run")
    cols = st.columns(4)
    selected_agents = []
    for i, agent in enumerate(AGENTS):
        with cols[i % 4]:
            if st.checkbox(f"{agent['icon']} {agent['name']}", value=True, key=f"chk_{agent['id']}"):
                selected_agents.append(agent)

    st.markdown('<br>', unsafe_allow_html=True)

    col_run, col_all = st.columns([1, 3])
    with col_run:
        run_all = st.button("⚡ Run Selected Agents", type="primary", use_container_width=True)

    if run_all and selected_agents:
        progress_bar = st.progress(0)
        status_text  = st.empty()

        for i, agent in enumerate(selected_agents):
            status_text.markdown(f'<div class="insight-item">🔄 Running <strong>{agent["name"]}</strong>...</div>', unsafe_allow_html=True)
            with st.spinner(f"Running {agent['name']}..."):
                result = run_agent(agent["id"], profile, df)
                st.session_state.agent_results[agent["id"]] = result
            progress_bar.progress((i + 1) / len(selected_agents))
            time.sleep(0.3)

        status_text.markdown('<div class="insight-item" style="border-color:#00D4AA;">✅ All agents completed!</div>', unsafe_allow_html=True)
        st.rerun()

    # Show results
    if st.session_state.agent_results:
        st.markdown("### 📋 Agent Reports")
        for agent in AGENTS:
            if agent["id"] in st.session_state.agent_results:
                with st.expander(f"{agent['icon']} {agent['name']}", expanded=False):
                    result = st.session_state.agent_results[agent["id"]]
                    st.markdown(result)

        col1, col2 = st.columns(2)
        with col1:
            if st.button("📊 View Dashboard →", type="primary", use_container_width=True):
                st.session_state.page = "dashboard"
                st.rerun()
        with col2:
            if st.button("🧠 Run ML Pipeline →", type="secondary", use_container_width=True):
                st.session_state.page = "ml"
                st.rerun()


def page_dashboard():
    st.markdown('<div class="section-title">📊 INTERACTIVE DASHBOARD</div>', unsafe_allow_html=True)

    df = st.session_state.cleaned_df or st.session_state.df
    if df is None:
        st.warning("Please upload a dataset first.")
        return

    numeric_cols     = df.select_dtypes(include=[np.number]).columns.tolist()
    categorical_cols = df.select_dtypes(include=["object", "category"]).columns.tolist()
    datetime_cols    = df.select_dtypes(include=["datetime64"]).columns.tolist()

    # ── Overview KPIs
    st.markdown("#### Dataset Overview")
    cols = st.columns(5)
    cols[0].metric("Total Rows",    f"{len(df):,}")
    cols[1].metric("Columns",        len(df.columns))
    cols[2].metric("Numeric",        len(numeric_cols))
    cols[3].metric("Categorical",    len(categorical_cols))
    cols[4].metric("Missing",        int(df.isnull().sum().sum()))

    if numeric_cols:
        st.divider()
        tab1, tab2, tab3, tab4, tab5 = st.tabs(["📈 Distributions", "🔗 Correlations", "📊 Categories", "📅 Time Series", "🔍 Custom"])

        PLOT_BG = "#07080F"
        PLOT_PAPER = "#07080F"
        PLOT_FONT  = "#E8E9F5"
        GRID_COL   = "#1E2235"

        def style_fig(fig):
            fig.update_layout(
                plot_bgcolor=PLOT_BG, paper_bgcolor=PLOT_PAPER,
                font_color=PLOT_FONT, font_family="DM Sans",
                legend=dict(bgcolor="#0C0E1A", bordercolor="#1E2235"),
                xaxis=dict(gridcolor=GRID_COL, zerolinecolor=GRID_COL),
                yaxis=dict(gridcolor=GRID_COL, zerolinecolor=GRID_COL),
                margin=dict(l=20, r=20, t=40, b=20),
            )
            return fig

        with tab1:
            selected_col = st.selectbox("Select column", numeric_cols)
            c1, c2 = st.columns(2)
            with c1:
                fig = px.histogram(df, x=selected_col, nbins=40, title=f"Distribution of {selected_col}",
                                   color_discrete_sequence=["#4361EE"])
                st.plotly_chart(style_fig(fig), use_container_width=True)
            with c2:
                fig = px.box(df, y=selected_col, title=f"Box Plot — {selected_col}",
                             color_discrete_sequence=["#7B2FFF"])
                st.plotly_chart(style_fig(fig), use_container_width=True)

            if len(numeric_cols) > 1:
                col2 = st.selectbox("Compare with", [c for c in numeric_cols if c != selected_col])
                fig = px.scatter(df, x=selected_col, y=col2, title=f"{selected_col} vs {col2}",
                                 color_discrete_sequence=["#00D4AA"], opacity=0.6)
                st.plotly_chart(style_fig(fig), use_container_width=True)

        with tab2:
            if len(numeric_cols) >= 2:
                corr = df[numeric_cols].corr().round(2)
                fig = px.imshow(corr, title="Correlation Matrix", color_continuous_scale="RdBu_r",
                                text_auto=True, aspect="auto")
                fig.update_layout(plot_bgcolor=PLOT_BG, paper_bgcolor=PLOT_PAPER, font_color=PLOT_FONT)
                st.plotly_chart(fig, use_container_width=True)

                # Top correlations
                corr_pairs = corr.abs().unstack().sort_values(ascending=False)
                corr_pairs = corr_pairs[corr_pairs < 1].drop_duplicates()
                st.markdown("**Top 10 Correlations:**")
                top_corr = pd.DataFrame(corr_pairs.head(10)).reset_index()
                top_corr.columns = ["Feature 1", "Feature 2", "Correlation"]
                st.dataframe(top_corr, use_container_width=True)
            else:
                st.info("Need at least 2 numeric columns for correlation analysis.")

        with tab3:
            if categorical_cols:
                cat_col = st.selectbox("Select categorical column", categorical_cols)
                val_counts = df[cat_col].value_counts().head(20)
                c1, c2 = st.columns(2)
                with c1:
                    fig = px.bar(x=val_counts.index, y=val_counts.values,
                                 title=f"Count by {cat_col}", color_discrete_sequence=["#4361EE"],
                                 labels={"x": cat_col, "y": "Count"})
                    st.plotly_chart(style_fig(fig), use_container_width=True)
                with c2:
                    fig = px.pie(values=val_counts.values, names=val_counts.index,
                                 title=f"Distribution of {cat_col}",
                                 color_discrete_sequence=px.colors.sequential.Plasma_r)
                    fig.update_layout(paper_bgcolor=PLOT_PAPER, font_color=PLOT_FONT)
                    st.plotly_chart(fig, use_container_width=True)

                if numeric_cols:
                    num_col = st.selectbox("Numeric column to compare", numeric_cols)
                    fig = px.box(df, x=cat_col, y=num_col, title=f"{num_col} by {cat_col}",
                                 color_discrete_sequence=["#7B2FFF"])
                    st.plotly_chart(style_fig(fig), use_container_width=True)
            else:
                st.info("No categorical columns found.")

        with tab4:
            if datetime_cols and numeric_cols:
                dt_col  = st.selectbox("Date column", datetime_cols)
                num_col = st.selectbox("Value column", numeric_cols, key="ts_num")
                freq    = st.selectbox("Frequency", ["D", "W", "M", "Q", "Y"])
                ts = df.set_index(dt_col)[num_col].resample(freq).sum().reset_index()
                fig = px.line(ts, x=dt_col, y=num_col, title=f"{num_col} over Time",
                              color_discrete_sequence=["#00D4AA"])
                fig.update_traces(fill="tozeroy", fillcolor="#00D4AA15")
                st.plotly_chart(style_fig(fig), use_container_width=True)
            else:
                st.info("No datetime columns detected. Parse date columns in the clean step.")

        with tab5:
            st.markdown("#### Build Your Own Chart")
            c1, c2, c3, c4 = st.columns(4)
            chart_type = c1.selectbox("Chart Type", ["Bar", "Line", "Scatter", "Histogram", "Box", "Pie"])
            x_col = c2.selectbox("X Axis", df.columns.tolist())
            y_col = c3.selectbox("Y Axis", numeric_cols) if numeric_cols else None
            color_col = c4.selectbox("Color by", ["None"] + categorical_cols)
            color_col = None if color_col == "None" else color_col

            if st.button("Generate Chart", type="primary"):
                try:
                    if chart_type == "Bar":
                        fig = px.bar(df, x=x_col, y=y_col, color=color_col, color_discrete_sequence=px.colors.sequential.Plasma_r)
                    elif chart_type == "Line":
                        fig = px.line(df, x=x_col, y=y_col, color=color_col)
                    elif chart_type == "Scatter":
                        fig = px.scatter(df, x=x_col, y=y_col, color=color_col, opacity=0.6)
                    elif chart_type == "Histogram":
                        fig = px.histogram(df, x=x_col, color=color_col)
                    elif chart_type == "Box":
                        fig = px.box(df, x=x_col, y=y_col, color=color_col)
                    elif chart_type == "Pie":
                        vc = df[x_col].value_counts().head(15)
                        fig = px.pie(values=vc.values, names=vc.index)

                    st.plotly_chart(style_fig(fig), use_container_width=True)
                except Exception as e:
                    st.error(f"Chart error: {e}")


def page_ml():
    st.markdown('<div class="section-title">🧠 ML PIPELINE</div>', unsafe_allow_html=True)

    df = st.session_state.cleaned_df or st.session_state.df
    if df is None:
        st.warning("Please upload a dataset first.")
        return

    numeric_cols = df.select_dtypes(include=[np.number]).columns.tolist()
    all_cols     = df.columns.tolist()

    if len(numeric_cols) < 2:
        st.warning("Need at least 2 numeric columns for ML.")
        return

    st.markdown("### Configure ML Task")
    c1, c2, c3 = st.columns(3)
    task       = c1.selectbox("Task Type", ["classification", "regression"])
    target_col = c2.selectbox("Target Column", all_cols)
    c3.markdown("<br>", unsafe_allow_html=True)

    if c3.button("⚡ Train Models", type="primary"):
        with st.spinner("Training 3 models and evaluating... This may take a moment."):
            try:
                results = run_ml_pipeline(df, target_col, task)
                st.session_state.ml_results = results
            except Exception as e:
                st.error(f"ML Error: {e}")
                return
        st.success("✅ Models trained successfully!")
        st.rerun()

    if st.session_state.ml_results:
        r = st.session_state.ml_results

        st.markdown("### 🏆 Results")
        if r["task"] == "classification":
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("Best Model",   r["best_model"])
            c2.metric("Accuracy",     f"{r['accuracy']:.2%}")
            c3.metric("F1 Score",     f"{r['f1']:.4f}")
            c4.metric("Test Samples", r["n_test"])
        else:
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("Best Model",   r["best_model"])
            c2.metric("R² Score",     f"{r['r2']:.4f}")
            c3.metric("MAE",          f"{r['mae']:,.2f}")
            c4.metric("Test Samples", r["n_test"])

        tab1, tab2, tab3 = st.tabs(["📊 Model Comparison", "🎯 Feature Importance", "📈 Predictions"])

        with tab1:
            scores_df = pd.DataFrame(r["model_scores"]).T.reset_index()
            scores_df.columns = ["Model"] + list(scores_df.columns[1:])
            metric_col = "accuracy" if r["task"] == "classification" else "r2"
            fig = px.bar(scores_df, x="Model", y=metric_col,
                         title="Model Comparison",
                         color="Model", color_discrete_sequence=["#4361EE","#7B2FFF","#00D4AA"])
            fig.update_layout(plot_bgcolor="#07080F", paper_bgcolor="#07080F",
                              font_color="#E8E9F5", showlegend=False)
            st.plotly_chart(fig, use_container_width=True)
            st.dataframe(scores_df, use_container_width=True)

        with tab2:
            if "feature_importance" in r:
                fi_df = pd.DataFrame(list(r["feature_importance"].items()), columns=["Feature", "Importance"])
                fi_df = fi_df.sort_values("Importance", ascending=True)
                fig = px.bar(fi_df, x="Importance", y="Feature", orientation="h",
                             title="Top Feature Importances",
                             color="Importance", color_continuous_scale="Blues")
                fig.update_layout(plot_bgcolor="#07080F", paper_bgcolor="#07080F", font_color="#E8E9F5")
                st.plotly_chart(fig, use_container_width=True)
            else:
                st.info("Feature importance not available for this model.")

        with tab3:
            y_test = r["y_test"][:50]
            y_pred = r["y_pred"][:50]
            if r["task"] == "regression":
                fig = go.Figure()
                fig.add_trace(go.Scatter(x=list(range(len(y_test))), y=y_test,
                                          mode="lines+markers", name="Actual",
                                          line=dict(color="#4361EE")))
                fig.add_trace(go.Scatter(x=list(range(len(y_pred))), y=y_pred,
                                          mode="lines+markers", name="Predicted",
                                          line=dict(color="#00D4AA", dash="dash")))
                fig.update_layout(title="Actual vs Predicted", plot_bgcolor="#07080F",
                                  paper_bgcolor="#07080F", font_color="#E8E9F5")
                st.plotly_chart(fig, use_container_width=True)
            else:
                pred_df = pd.DataFrame({"Actual": [str(x) for x in y_test],
                                         "Predicted": [str(x) for x in y_pred]})
                correct = sum(1 for a, p in zip(y_test, y_pred) if a == p)
                st.metric("Correct Predictions (sample)", f"{correct}/{len(y_test)}")
                st.dataframe(pred_df.head(20), use_container_width=True)

        if st.button("📄 Generate Reports →", type="primary"):
            st.session_state.page = "reports"
            st.rerun()


def page_reports():
    st.markdown('<div class="section-title">📄 REPORT GENERATOR</div>', unsafe_allow_html=True)

    df = st.session_state.cleaned_df or st.session_state.df
    if df is None:
        st.warning("Please upload a dataset first.")
        return

    profile = get_data_profile(df)

    if not st.session_state.agent_results:
        st.warning("Run AI Agents first for richer reports.")

    st.markdown("### Generate Your Reports")
    c1, c2 = st.columns(2)

    with c1:
        st.markdown('<div class="nabbr-card">', unsafe_allow_html=True)
        st.markdown("#### 📕 PDF Report")
        st.markdown("Complete analytics report with all AI agent insights, ML results, and data profile — ready to share with stakeholders.")
        if st.button("⬇️ Download PDF Report", type="primary", use_container_width=True):
            with st.spinner("Generating PDF..."):
                pdf_bytes = generate_pdf_report(
                    st.session_state.filename or "dataset",
                    profile,
                    st.session_state.agent_results,
                    st.session_state.ml_results,
                )
            st.download_button(
                label="📕 Download PDF",
                data=pdf_bytes,
                file_name=f"NABBR_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )
        st.markdown('</div>', unsafe_allow_html=True)

    with c2:
        st.markdown('<div class="nabbr-card">', unsafe_allow_html=True)
        st.markdown("#### 📗 Excel Report")
        st.markdown("Structured Excel workbook with Executive Summary, Raw Data (1000 rows), Statistics, and ML Results across 4 sheets.")
        if st.button("⬇️ Download Excel Report", type="secondary", use_container_width=True):
            with st.spinner("Generating Excel..."):
                excel_bytes = generate_excel_report(
                    df, profile,
                    st.session_state.agent_results,
                    st.session_state.ml_results,
                )
            st.download_button(
                label="📗 Download Excel",
                data=excel_bytes,
                file_name=f"NABBR_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        st.markdown('</div>', unsafe_allow_html=True)

    # Report preview
    if st.session_state.agent_results:
        st.markdown("### 📋 Report Preview")
        for agent in AGENTS:
            if agent["id"] in st.session_state.agent_results:
                with st.expander(f"{agent['icon']} {agent['name']} Report"):
                    st.markdown(st.session_state.agent_results[agent["id"]])

    # Clean data download
    st.markdown("### ⬇️ Download Cleaned Data")
    if st.session_state.cleaned_df is not None:
        csv = st.session_state.cleaned_df.to_csv(index=False).encode()
        st.download_button("📥 Download Cleaned CSV", data=csv,
                           file_name=f"cleaned_{st.session_state.filename or 'data.csv'}",
                           mime="text/csv")


def page_history():
    st.markdown('<div class="section-title">📜 SESSION HISTORY</div>', unsafe_allow_html=True)

    workspaces = get_workspaces()
    if not workspaces:
        st.info("No workspaces yet. Upload a dataset to create your first session.")
        return

    for ws in workspaces:
        sessions = get_sessions(ws["id"])
        if not sessions:
            continue

        st.markdown(f"#### 📁 {ws['name']}")
        for session in sessions:
            created = session.get("created_at", "")[:16].replace("T", " ")
            col1, col2, col3, col4 = st.columns([3, 1, 1, 1])
            col1.markdown(f"**{session['name']}**  \n<span style='font-size:11px;color:#7880A0;'>{session.get('filename','')}</span>", unsafe_allow_html=True)
            col2.markdown(f'<span class="stat-pill">{session.get("row_count",0):,} rows</span>', unsafe_allow_html=True)
            col3.markdown(f'<span class="stat-pill">{session.get("col_count",0)} cols</span>', unsafe_allow_html=True)
            col4.markdown(f'<span style="font-size:11px;color:#3A4060;font-family:\'JetBrains Mono\',monospace;">{created}</span>', unsafe_allow_html=True)
            st.divider()


# ── MAIN ROUTER ───────────────────────────────────────────────────────────────
def main():
    render_sidebar()

    if not st.session_state.user:
        page_login()
        return

    page = st.session_state.page
    if page == "home":       page_home()
    elif page == "upload":   page_upload()
    elif page == "agents":   page_agents()
    elif page == "dashboard": page_dashboard()
    elif page == "ml":       page_ml()
    elif page == "reports":  page_reports()
    elif page == "history":  page_history()
    else:                    page_home()

if __name__ == "__main__":
    main()
